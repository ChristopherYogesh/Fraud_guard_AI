import os
from datetime import datetime
from pathlib import Path

class ReportGenerator:
    """Generates PDF and Excel reports for Fraud Guard AI security operations."""
    
    @staticmethod
    def generate_pdf_report(report_title, period, incidents, summary_stats, output_path):
        """Generates a professional executive PDF report using ReportLab."""
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib import colors
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
            
            doc = SimpleDocTemplate(
                str(output_path),
                pagesize=letter,
                rightMargin=36,
                leftMargin=36,
                topMargin=36,
                bottomMargin=36
            )
            
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'ReportTitle',
                parent=styles['Heading1'],
                fontSize=22,
                leading=26,
                textColor=colors.HexColor('#0F172A'),
                fontName='Helvetica-Bold'
            )
            subtitle_style = ParagraphStyle(
                'ReportSubTitle',
                parent=styles['Normal'],
                fontSize=11,
                leading=14,
                textColor=colors.HexColor('#475569')
            )
            heading2_style = ParagraphStyle(
                'SectionHeading',
                parent=styles['Heading2'],
                fontSize=14,
                leading=18,
                textColor=colors.HexColor('#1E293B'),
                spaceBefore=14,
                spaceAfter=8
            )
            cell_style = ParagraphStyle(
                'TableCell',
                parent=styles['Normal'],
                fontSize=9,
                leading=11,
                textColor=colors.HexColor('#1E293B')
            )
            
            story = []
            
            # Header
            story.append(Paragraph("FRAUD GUARD AI — SECURITY INTELLIGENCE REPORT", title_style))
            story.append(Paragraph(f"<b>Report:</b> {report_title} | <b>Period:</b> {period} | <b>Generated:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}", subtitle_style))
            story.append(Spacer(1, 10))
            story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0284C7'), spaceBefore=5, spaceAfter=15))
            
            # Executive Summary KPI Table
            story.append(Paragraph("Executive Summary & Operational KPIs", heading2_style))
            summary_data = [
                ['Metric', 'Value', 'Status / Benchmark'],
                ['Total Detections Logged', str(summary_stats.get('total', 0)), 'Across all monitored zones'],
                ['Critical / High Threats', str(summary_stats.get('critical', 0) + summary_stats.get('high', 0)), 'Required immediate alert dispatch'],
                ['Confirmed Incidents', str(summary_stats.get('confirmed', 0)), 'Verified security interventions'],
                ['Avg AI Confidence', f"{summary_stats.get('avg_conf', 94.5):.1f}%", 'YOLOv8 + CNN fused inference']
            ]
            summary_table = Table(summary_data, colWidths=[180, 100, 260])
            summary_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F8FAFC'), colors.white])
            ]))
            story.append(summary_table)
            story.append(Spacer(1, 15))
            
            # Incidents Log Table
            story.append(Paragraph("Detailed Incident & Suspicion Log", heading2_style))
            incident_data = [['ID', 'Timestamp', 'Camera Zone', 'Classification', 'Confidence', 'Risk Level', 'Status']]
            
            for inc in incidents[:40]:  # Limit top 40 for clean PDF layout
                incident_data.append([
                    f"#{inc.id}",
                    inc.timestamp.strftime('%Y-%m-%d %H:%M'),
                    inc.camera,
                    inc.object_class.title(),
                    f"{int(inc.confidence * 100)}%",
                    inc.risk_level,
                    inc.status
                ])
                
            if len(incidents) == 0:
                incident_data.append(['—', 'No incidents recorded in this period', '', '', '', '', ''])
                
            inc_table = Table(incident_data, colWidths=[40, 100, 130, 90, 60, 60, 60])
            inc_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 8),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('ALIGN', (2, 1), (2, -1), 'LEFT'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F1F5F9'), colors.white])
            ]))
            story.append(inc_table)
            
            # Build PDF
            doc.build(story)
            return True
        except Exception as e:
            # Fallback simple text report if ReportLab fails
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(f"FRAUD GUARD AI REPORT - {report_title}\nPeriod: {period}\n\nSummary:\n{summary_stats}\n\nIncidents:\n")
                for inc in incidents:
                    f.write(f"#{inc.id} | {inc.timestamp} | {inc.camera} | {inc.object_class} | {inc.confidence} | {inc.risk_level}\n")
            return True
            
    @staticmethod
    def generate_excel_report(report_title, period, incidents, summary_stats, output_path):
        """Generates an auditable Excel workbook using openpyxl."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Security Report"
            
            # Title
            ws.merge_cells('A1:G1')
            ws['A1'] = "FRAUD GUARD AI — SECURITY INCIDENT REPORT"
            ws['A1'].font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
            ws['A1'].fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
            ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 35
            
            ws.merge_cells('A2:G2')
            ws['A2'] = f"Report: {report_title} | Period: {period} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            ws['A2'].font = Font(name='Calibri', size=10, italic=True)
            ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[2].height = 20
            
            # Table Headers
            headers = ['Incident ID', 'Timestamp', 'Camera / Zone', 'Classification', 'Confidence', 'Risk Level', 'Investigation Status']
            ws.append([])  # blank row 3
            ws.append(headers)  # row 4
            
            header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=4, column=col_num)
                cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
                
            ws.row_dimensions[4].height = 24
            
            # Rows
            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )
            
            for row_idx, inc in enumerate(incidents, 5):
                ws.append([
                    f"INC-{inc.id:04d}",
                    inc.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
                    inc.camera,
                    inc.object_class.title(),
                    f"{int(inc.confidence * 100)}%",
                    inc.risk_level,
                    inc.status
                ])
                for col_num in range(1, 8):
                    c = ws.cell(row=row_idx, column=col_num)
                    c.border = thin_border
                    c.alignment = Alignment(horizontal='center' if col_num in [1, 5, 6, 7] else 'left')
                    
            # Column auto-width
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
                
            wb.save(output_path)
            return True
        except Exception as e:
            # Fallback CSV
            import csv
            with open(output_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                writer.writerow(['ID', 'Timestamp', 'Camera', 'Class', 'Confidence', 'Risk Level', 'Status'])
                for inc in incidents:
                    writer.writerow([inc.id, inc.timestamp, inc.camera, inc.object_class, inc.confidence, inc.risk_level, inc.status])
            return True
